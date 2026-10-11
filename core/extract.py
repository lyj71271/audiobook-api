"""포맷별 텍스트 추출. 스캔본 PDF는 예외로 중단."""
import os


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".txt", ".md", ".text"):
        return _from_txt(path)
    if ext == ".pdf":
        return _from_pdf(path)
    if ext == ".epub":
        # 챕터 구분이 필요하면 extract_chapters() 사용
        return "\n\n".join(extract_chapters(path))
    if ext == ".docx":
        return _from_docx(path)
    raise ValueError("지원하지 않는 형식: %s (txt/pdf/epub/docx만 지원)" % ext)


def _read_auto(path: str) -> str:
    # utf-8-sig 우선, 실패 시 charset-normalizer
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            return f.read()
    except UnicodeDecodeError:
        from charset_normalizer import from_path
        return str(from_path(path).best())


def _from_txt(path: str) -> str:
    return _read_auto(path)


def _from_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    text = "\n".join(parts)
    if len(text.strip()) < 100:
        raise ValueError("스캔본 PDF 의심 — 텍스트가 100자 미만. OCR이 필요합니다.")
    return text


def extract_chapters(path: str):
    """epub → 챕터 리스트. epub 아니면 단일 챕터 리스트 반환."""
    ext = os.path.splitext(path)[1].lower()
    if ext != ".epub":
        return [extract_text(path)]
    from ebooklib import epub
    from bs4 import BeautifulSoup
    book = epub.read_epub(path)
    chapters = []
    for item in book.get_items():
        # ITEM_DOCUMENT == 9
        try:
            if item.get_type() != 9:
                continue
        except Exception:
            continue
        soup = BeautifulSoup(item.get_content(), "html.parser")
        t = soup.get_text("\n").strip()
        if len(t) >= 50:  # 목차/빈 페이지 제외
            chapters.append(t)
    if not chapters:
        raise ValueError("epub에서 본문을 찾지 못했습니다.")
    return chapters


def _from_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)
